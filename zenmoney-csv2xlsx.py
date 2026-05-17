from pathlib import Path
 
import pandas as pd
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.utils import get_column_letter


AMOUNT_COLUMNS = [
    "amount",
    "amount_abs",
    "counterparty_amount",
    "source_outcome",
    "source_income",
]


def is_filled(value) -> bool:
    """Проверяет, что значение не пустое."""
    return pd.notna(value) and str(value).strip() != ""


def clean_value_for_excel(value):
    """
    Удаляет символы, которые запрещены внутри XLSX/XML.

    CSV такие символы может пережить, а Excel-файл через openpyxl — нет.
    Также ограничивает длину текста лимитом Excel: 32 767 символов на ячейку.
    """
    if isinstance(value, str):
        value = ILLEGAL_CHARACTERS_RE.sub("", value)
        return value[:32767]

    return value


def clean_dataframe_for_excel(df: pd.DataFrame) -> pd.DataFrame:
    """
    Создает копию DataFrame и чистит только текстовые колонки.
    Числовые колонки остаются числами.
    """
    cleaned = df.copy()

    text_columns = cleaned.select_dtypes(include=["object"]).columns

    for col in text_columns:
        cleaned[col] = cleaned[col].map(clean_value_for_excel)

    return cleaned


def parse_amount_series(series: pd.Series) -> pd.Series:
    """
    Преобразует колонку с суммами в числа.

    Поддерживает варианты:
    - 1234.56
    - 1234,56
    - "1234.56"
    - "1 234,56"
    """
    prepared = (
        series
        .astype(str)
        .str.replace("\u00a0", "", regex=False)
        .str.replace(" ", "", regex=False)
        .str.replace(",", ".", regex=False)
        .str.strip()
    )

    prepared = prepared.replace({
        "": "0",
        "nan": "0",
        "None": "0",
        "NaN": "0",
    })

    return pd.to_numeric(prepared, errors="coerce").fillna(0)


def load_zenmoney_csv(input_file: str | Path) -> pd.DataFrame:
    """
    Загружает CSV-файл ZenMoney.

    Входной файл читается как UTF-8 без BOM.
    Если BOM фактически присутствует, он удаляется из названия первой колонки.
    """
    input_file = Path(input_file)

    df = pd.read_csv(
        input_file,
        encoding="utf-8",
        sep=None,
        engine="python",
        dtype=str,
        keep_default_na=False,
    )

    df.columns = [
        str(col).replace("\ufeff", "").strip()
        for col in df.columns
    ]

    return df


def transform_zenmoney_csv(input_file: str | Path) -> pd.DataFrame:
    """
    Преобразует выгрузку ZenMoney в плоский файл движений по счетам.

    Логика:
    - обычный расход -> одна строка с amount < 0;
    - обычный доход -> одна строка с amount > 0;
    - перевод между счетами -> две строки:
      1. transfer_out — расход со счета списания;
      2. transfer_in — доход на счет зачисления.
    """
    df = load_zenmoney_csv(input_file)

    required_columns = [
        "date",
        "categoryName",
        "payee",
        "comment",
        "outcomeAccountName",
        "outcome",
        "outcomeCurrencyShortTitle",
        "incomeAccountName",
        "income",
        "incomeCurrencyShortTitle",
        "createdDate",
        "changedDate",
        "qrCode",
    ]

    missing_columns = [col for col in required_columns if col not in df.columns]

    if missing_columns:
        raise ValueError(
            "В исходном файле отсутствуют обязательные колонки: "
            + ", ".join(missing_columns)
            + "\nФактические колонки: "
            + ", ".join(df.columns)
        )

    df["outcome"] = parse_amount_series(df["outcome"])
    df["income"] = parse_amount_series(df["income"])

    rows = []

    for idx, row in df.iterrows():
        original_file_row = idx + 2
        original_operation_id = idx + 1

        outcome_account = row["outcomeAccountName"]
        outcome_amount = row["outcome"]
        outcome_currency = row["outcomeCurrencyShortTitle"]

        income_account = row["incomeAccountName"]
        income_amount = row["income"]
        income_currency = row["incomeCurrencyShortTitle"]

        base_fields = {
            "original_operation_id": original_operation_id,
            "original_file_row": original_file_row,
            "date": row["date"],
            "category_name": row["categoryName"],
            "payee": row["payee"],
            "comment": row["comment"],
            "created_date": row["createdDate"],
            "changed_date": row["changedDate"],
            "qr_code": row["qrCode"],
            "source_outcome_account_name": outcome_account,
            "source_outcome": outcome_amount,
            "source_outcome_currency": outcome_currency,
            "source_income_account_name": income_account,
            "source_income": income_amount,
            "source_income_currency": income_currency,
        }

        # Перевод между счетами:
        # одна исходная строка превращается в две строки.
        if outcome_amount > 0 and income_amount > 0:
            if is_filled(outcome_account):
                rows.append({
                    **base_fields,
                    "movement_id": f"{original_operation_id}_out",
                    "operation_type": "transfer",
                    "movement_type": "transfer_out",
                    "direction": "out",
                    "account_name": outcome_account,
                    "account_currency": outcome_currency,
                    "amount": -outcome_amount,
                    "amount_abs": outcome_amount,
                    "counterparty_account_name": income_account,
                    "counterparty_currency": income_currency,
                    "counterparty_amount": income_amount,
                })

            if is_filled(income_account):
                rows.append({
                    **base_fields,
                    "movement_id": f"{original_operation_id}_in",
                    "operation_type": "transfer",
                    "movement_type": "transfer_in",
                    "direction": "in",
                    "account_name": income_account,
                    "account_currency": income_currency,
                    "amount": income_amount,
                    "amount_abs": income_amount,
                    "counterparty_account_name": outcome_account,
                    "counterparty_currency": outcome_currency,
                    "counterparty_amount": outcome_amount,
                })

        # Обычный расход.
        elif outcome_amount > 0:
            if is_filled(outcome_account):
                rows.append({
                    **base_fields,
                    "movement_id": f"{original_operation_id}_expense",
                    "operation_type": "expense",
                    "movement_type": "expense",
                    "direction": "out",
                    "account_name": outcome_account,
                    "account_currency": outcome_currency,
                    "amount": -outcome_amount,
                    "amount_abs": outcome_amount,
                    "counterparty_account_name": "",
                    "counterparty_currency": "",
                    "counterparty_amount": pd.NA,
                })

        # Обычный доход.
        elif income_amount > 0:
            if is_filled(income_account):
                rows.append({
                    **base_fields,
                    "movement_id": f"{original_operation_id}_income",
                    "operation_type": "income",
                    "movement_type": "income",
                    "direction": "in",
                    "account_name": income_account,
                    "account_currency": income_currency,
                    "amount": income_amount,
                    "amount_abs": income_amount,
                    "counterparty_account_name": "",
                    "counterparty_currency": "",
                    "counterparty_amount": pd.NA,
                })

        # Строки без движения денег не включаем.
        else:
            continue

    result = pd.DataFrame(rows)

    columns_order = [
        "movement_id",
        "original_operation_id",
        "original_file_row",
        "date",
        "operation_type",
        "movement_type",
        "direction",
        "account_name",
        "account_currency",
        "amount",
        "amount_abs",
        "counterparty_account_name",
        "counterparty_currency",
        "counterparty_amount",
        "category_name",
        "payee",
        "comment",
        "created_date",
        "changed_date",
        "qr_code",
        "source_outcome_account_name",
        "source_outcome",
        "source_outcome_currency",
        "source_income_account_name",
        "source_income",
        "source_income_currency",
    ]

    result = result[columns_order]

    for col in AMOUNT_COLUMNS:
        result[col] = pd.to_numeric(result[col], errors="coerce")

    return result


def save_result_to_csv(result: pd.DataFrame, output_csv_file: str | Path) -> None:
    """
    Сохраняет результат в CSV:
    - разделитель колонок: ;
    - десятичный разделитель: ,
    - кодировка: UTF-8 без BOM.
    """
    output_csv_file = Path(output_csv_file)

    result.to_csv(
        output_csv_file,
        index=False,
        sep=";",
        decimal=",",
        encoding="utf-8",
        float_format="%.2f",
        na_rep="",
    )


def autofit_columns(worksheet, max_width: int = 60, sample_rows: int = 1000) -> None:
    """
    Безопасно подбирает ширину колонок.

    Ограничиваем количество строк для анализа, чтобы Excel не тормозил
    и чтобы не было нестабильного поведения на больших файлах.
    """
    max_row = min(worksheet.max_row, sample_rows)

    for col_idx in range(1, worksheet.max_column + 1):
        column_letter = get_column_letter(col_idx)
        max_length = 0

        for row_idx in range(1, max_row + 1):
            value = worksheet.cell(row=row_idx, column=col_idx).value

            if value is not None:
                max_length = max(max_length, len(str(value)))

        worksheet.column_dimensions[column_letter].width = min(max_length + 2, max_width)


def apply_number_format(worksheet, column_names: list[str], number_format: str) -> None:
    """
    Применяет числовой формат к колонкам по их названиям.
    """
    headers = {
        worksheet.cell(row=1, column=col_idx).value: col_idx
        for col_idx in range(1, worksheet.max_column + 1)
    }

    for column_name in column_names:
        col_idx = headers.get(column_name)

        if col_idx is None:
            continue

        for row_idx in range(2, worksheet.max_row + 1):
            worksheet.cell(row=row_idx, column=col_idx).number_format = number_format


def save_result_to_excel(result: pd.DataFrame, output_excel_file: str | Path) -> None:
    """
    Сохраняет результат в Excel-файл.

    Исправления по сравнению с прежней версией:
    - удаляются запрещенные для XLSX управляющие символы;
    - Excel пишется сначала во временный файл;
    - старый файл заменяется только после успешной записи нового;
    - автоширина колонок ограничена первыми 1000 строками;
    - денежные поля сохраняются как числа, а не как текст.
    """
    output_excel_file = Path(output_excel_file)
    temp_excel_file = output_excel_file.with_suffix(".tmp.xlsx")

    if temp_excel_file.exists():
        temp_excel_file.unlink()

    excel_result = clean_dataframe_for_excel(result)

    balance_by_account = (
        excel_result
        .groupby(["account_name", "account_currency"], as_index=False)
        .agg(
            balance=("amount", "sum"),
            movement_count=("movement_id", "count"),
        )
        .sort_values(["account_name", "account_currency"])
    )

    try:
        with pd.ExcelWriter(temp_excel_file, engine="openpyxl", mode="w") as writer:
            excel_result.to_excel(
                writer,
                sheet_name="movements",
                index=False,
            )

            balance_by_account.to_excel(
                writer,
                sheet_name="balance_by_account",
                index=False,
            )

            movements_sheet = writer.sheets["movements"]
            balance_sheet = writer.sheets["balance_by_account"]

            for worksheet in [movements_sheet, balance_sheet]:
                worksheet.freeze_panes = "A2"
                worksheet.auto_filter.ref = worksheet.dimensions
                autofit_columns(worksheet)

            apply_number_format(
                movements_sheet,
                column_names=[
                    "amount",
                    "amount_abs",
                    "counterparty_amount",
                    "source_outcome",
                    "source_income",
                ],
                number_format="#,##0.00",
            )

            apply_number_format(
                balance_sheet,
                column_names=["balance"],
                number_format="#,##0.00",
            )

        if output_excel_file.exists():
            output_excel_file.unlink()

        temp_excel_file.replace(output_excel_file)

    except PermissionError as error:
        if temp_excel_file.exists():
            temp_excel_file.unlink()

        raise PermissionError(
            f"Не удалось записать Excel-файл '{output_excel_file}'. "
            f"Скорее всего, он открыт в Excel или другой программе. "
            f"Закрой файл и запусти скрипт еще раз."
        ) from error

    except Exception:
        if temp_excel_file.exists():
            temp_excel_file.unlink()

        raise


def print_processing_report(input_file: str | Path, result: pd.DataFrame) -> None:
    """
    Печатает контрольную статистику обработки.
    """
    source_df = load_zenmoney_csv(input_file)

    source_df["outcome"] = parse_amount_series(source_df["outcome"])
    source_df["income"] = parse_amount_series(source_df["income"])

    source_rows = len(source_df)
    transfer_rows = len(source_df[(source_df["outcome"] > 0) & (source_df["income"] > 0)])
    zero_rows = len(source_df[(source_df["outcome"] == 0) & (source_df["income"] == 0)])
    expected_rows = source_rows + transfer_rows - zero_rows

    print("Готово")
    print(f"Исходных строк операций: {source_rows}")
    print(f"Переводов между счетами: {transfer_rows}")
    print(f"Строк без движения денег: {zero_rows}")
    print(f"Ожидаемое количество строк в выходном файле: {expected_rows}")
    print(f"Фактическое количество строк в выходном файле: {len(result)}")


def main() -> None:
    input_file = Path("input.csv")

    output_csv_file = Path("zenmoney_bi_flat_v2.csv")
    output_excel_file = Path("zenmoney_bi_flat_v2.xlsx")

    result = transform_zenmoney_csv(input_file)

    save_result_to_csv(result, output_csv_file)
    save_result_to_excel(result, output_excel_file)

    print_processing_report(input_file, result)

    print(f"CSV-файл: {output_csv_file}")
    print(f"Excel-файл: {output_excel_file}")


if __name__ == "__main__":
    main()